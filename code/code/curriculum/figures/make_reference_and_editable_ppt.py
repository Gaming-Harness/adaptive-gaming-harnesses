"""Package the exact generated reference and editable redraw into one PPTX."""

from copy import deepcopy
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.util import Inches


HERE = Path(__file__).resolve().parent
REFERENCE = HERE / "probing_harness_for_gaming_agents_v1.png"
EDITABLE = HERE / "probing_harness_for_gaming_agents_editable.pptx"
OUTPUT = HERE / "probing_harness_reference_and_editable.pptx"


def main() -> None:
    src = Presentation(EDITABLE)
    prs = Presentation()
    prs.slide_width = src.slide_width
    prs.slide_height = src.slide_height

    # Slide 1: pixel-identical reference image, fitted without distortion.
    ref = prs.slides.add_slide(prs.slide_layouts[6])
    with Image.open(REFERENCE) as im:
        aspect = im.width / im.height
    slide_w = prs.slide_width / Inches(1)
    slide_h = prs.slide_height / Inches(1)
    if slide_w / slide_h >= aspect:
        h = slide_h
        w = h * aspect
        x, y = (slide_w - w) / 2, 0
    else:
        w = slide_w
        h = w / aspect
        x, y = 0, (slide_h - h) / 2
    ref.shapes.add_picture(str(REFERENCE), Inches(x), Inches(y), Inches(w), Inches(h))

    # Slide 2: every panel, arrow, line and label remains independently editable.
    editable = prs.slides.add_slide(prs.slide_layouts[6])
    for shape in src.slides[0].shapes:
        editable.shapes._spTree.insert_element_before(
            deepcopy(shape.element), "p:extLst"
        )

    prs.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
