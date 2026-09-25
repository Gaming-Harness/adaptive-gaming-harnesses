"""Create an editable PowerPoint method diagram for Probing Harness."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


OUT = Path(__file__).with_name("probing_harness_for_gaming_agents_editable.pptx")

BLUE = "1967E8"
BLUE_DARK = "1053B4"
BLUE_BG = "F1F6FF"
GREEN = "17853A"
GREEN_DARK = "11732E"
GREEN_BG = "F2FAEE"
ORANGE = "E57A00"
ORANGE_BG = "FFF7E3"
RED = "C91735"
GRAY = "68707E"
LIGHT_GRAY = "F7F8FA"
BLACK = "202124"
WHITE = "FFFFFF"


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color)


def add_box(slide, x, y, w, h, *, fill, line, radius=True, width=1.5):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    s = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    s.fill.solid()
    s.fill.fore_color.rgb = rgb(fill)
    s.line.color.rgb = rgb(line)
    s.line.width = Pt(width)
    return s


def add_text(slide, text, x, y, w, h, *, size=14, color=BLACK, bold=False,
             align=PP_ALIGN.CENTER, font="Aptos", valign=MSO_ANCHOR.MIDDLE):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = Inches(0.02)
    tf.margin_top = tf.margin_bottom = Inches(0.01)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = rgb(color)
    return tb


def add_line(slide, x1, y1, x2, y2, *, color=BLACK, width=1.7,
             dash=False, begin=False, end=True):
    ln = slide.shapes.add_connector(
        MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2)
    )
    ln.line.color.rgb = rgb(color)
    ln.line.width = Pt(width)
    if dash:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    if begin:
        ln.line.begin_arrowhead = True
    if end:
        ln.line.end_arrowhead = True
    return ln


def add_node(slide, label, x, y, w, h, *, accent=BLUE, icon=""):
    add_box(slide, x, y, w, h, fill=WHITE, line=accent, radius=True, width=1.3)
    if icon:
        add_text(slide, icon, x + 0.05, y + 0.05, w - 0.1, h * 0.42,
                 size=20, color=accent, bold=True)
        ty, th = y + h * 0.43, h * 0.52
    else:
        ty, th = y + 0.03, h - 0.06
    add_text(slide, label, x + 0.05, ty, w - 0.1, th, size=10.5, bold=True)


def build() -> None:
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(WHITE)

    add_text(slide, "Probing Harness for Gaming Agents", 0.15, 0.08, 13.0, 0.42,
             size=25, color=BLUE, bold=True)

    # Current harness.
    add_box(slide, 0.18, 1.86, 2.25, 3.34, fill=LIGHT_GRAY, line=GRAY)
    add_text(slide, "Current Gaming Harness  Hₜ", 0.28, 1.98, 2.05, 0.35,
             size=14, bold=True)
    add_text(slide, "task knowledge · probes · memory\naction middleware", 0.35, 2.34, 1.9, 0.44,
             size=9.3, color=GRAY)
    add_box(slide, 0.55, 2.94, 1.5, 1.15, fill=WHITE, line=BLACK, radius=True)
    add_text(slide, "🎮", 0.63, 3.05, 0.42, 0.35, size=19, color=BLUE)
    add_text(slide, "Prompt", 1.08, 3.02, 0.75, 0.24, size=9, bold=True)
    add_text(slide, "Probe", 1.08, 3.31, 0.75, 0.24, size=9, bold=True)
    add_text(slide, "Memory", 1.08, 3.60, 0.75, 0.24, size=9, bold=True)
    add_line(slide, 1.30, 4.08, 1.30, 4.34, color=BLACK, end=True)
    frozen_left = add_box(slide, 0.42, 4.35, 1.78, 0.55, fill=WHITE, line=GRAY,
                          radius=True, width=1.1)
    frozen_left.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_text(slide, "Frozen Gaming Model  M", 0.48, 4.44, 1.66, 0.32,
             size=11.2, bold=True)

    # Probing stage.
    add_box(slide, 3.58, 0.55, 6.12, 2.16, fill=BLUE_BG, line=BLUE_DARK)
    add_text(slide, "Gaming Weakness Probing", 3.80, 0.67, 5.68, 0.38,
             size=20, color=BLUE, bold=True)
    add_node(slide, "Run Hₜ on\nGame Tasks", 3.82, 1.17, 1.05, 1.10,
             accent=BLUE_DARK, icon="🎮")
    add_line(slide, 4.87, 1.72, 5.12, 1.72, color=BLUE, end=True)
    add_node(slide, "Perception–Action\nTraces", 5.12, 1.17, 1.25, 1.10,
             accent=BLUE_DARK, icon="●—●—●")
    add_line(slide, 6.37, 1.72, 6.62, 1.72, color=BLUE, end=True)
    add_box(slide, 6.62, 1.13, 1.58, 1.20, fill=WHITE, line=BLUE,
            radius=True, width=1.1)
    add_text(slide, "ACTIVE PROBES", 6.73, 1.20, 1.36, 0.20,
             size=9, color=BLUE, bold=True)
    probes = ["Target Alignment", "Reachability", "Attack Persistence",
              "Tool / Inventory", "Success Verification"]
    for i, t in enumerate(probes):
        add_text(slide, "✓  " + t, 6.77, 1.43 + i * 0.16, 1.30, 0.16,
                 size=7.3, align=PP_ALIGN.LEFT, color=BLACK)
    add_line(slide, 8.20, 1.72, 8.42, 1.72, color=BLUE, end=True)
    add_node(slide, "Clustered\nFailure Modes", 8.42, 1.17, 1.02, 1.10,
             accent=RED, icon="● ●  ●")
    add_text(slide, "Only Turns · Stalls after Alignment\nWrong Tool · No Completion Check",
             7.88, 2.35, 1.63, 0.26, size=6.6, color=RED)
    add_line(slide, 2.43, 2.32, 3.58, 2.32, color=BLACK, end=True)
    add_line(slide, 6.64, 2.71, 6.64, 3.02, color=BLUE, end=True)
    add_text(slide, "Failures + Probe Evidence", 6.77, 2.74, 1.75, 0.30,
             size=11, color=BLUE, bold=True, align=PP_ALIGN.LEFT)

    # Proposal stage.
    add_box(slide, 3.58, 3.05, 6.12, 1.90, fill=GREEN_BG, line=GREEN_DARK)
    add_text(slide, "Harness Proposal", 3.80, 3.15, 5.68, 0.37,
             size=20, color=GREEN_DARK, bold=True)
    add_text(slide, "Selected\nFailure Modes", 3.82, 3.73, 1.15, 0.58,
             size=11.5, bold=True)
    for i, t in enumerate(["only turns", "stalls", "wrong tool"]):
        add_text(slide, "W%d: %s" % (i + 1, t), 3.82, 4.31 + i * 0.16,
                 1.12, 0.15, size=7.5, align=PP_ALIGN.LEFT)
    add_line(slide, 4.98, 4.05, 5.22, 4.05, color=GREEN, end=True)
    add_box(slide, 5.22, 3.66, 1.32, 0.96, fill=WHITE, line=GRAY, radius=True)
    add_text(slide, "Hₜ as Proposer", 5.33, 3.75, 1.10, 0.28, size=11, bold=True)
    frozen_mid = add_box(slide, 5.38, 4.10, 1.00, 0.35, fill=LIGHT_GRAY,
                         line=GRAY, radius=True, width=0.9)
    frozen_mid.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_text(slide, "Frozen M", 5.45, 4.15, 0.86, 0.20, size=9.4, bold=True)
    add_line(slide, 6.54, 4.05, 6.82, 4.05, color=GREEN, end=True)
    add_text(slide, "Proposed Harness Edits", 6.85, 3.55, 2.56, 0.28,
             size=12.2, color=GREEN_DARK, bold=True)
    edits = ["Camera Alignment Controller", "Hold-Attack Policy",
             "Recovery / Loop Breaker", "Tool Selection Rule",
             "Success Verifier", "Task-Knowledge Retrieval"]
    for i, t in enumerate(edits):
        col, row = i // 3, i % 3
        add_text(slide, "＋ " + t, 6.86 + col * 1.28, 3.88 + row * 0.27,
                 1.22, 0.24, size=7.8, color=GREEN_DARK,
                 align=PP_ALIGN.LEFT, bold=True)
    add_text(slide, "UPDATE HARNESS ONLY — MODEL M REMAINS FROZEN",
             6.88, 4.68, 2.45, 0.18, size=7.4, color=GREEN_DARK, bold=True)
    add_line(slide, 6.64, 4.95, 6.64, 5.25, color=GREEN, end=True)
    add_text(slide, "Candidate Harnesses  H′ₜ", 6.76, 4.98, 1.92, 0.28,
             size=11, color=GREEN, bold=True, align=PP_ALIGN.LEFT)

    # Validation stage.
    add_box(slide, 3.70, 5.27, 5.88, 1.93, fill=ORANGE_BG, line=ORANGE)
    add_text(slide, "Paired Proposal Validation", 3.93, 5.37, 5.42, 0.36,
             size=19, color=ORANGE, bold=True)
    add_box(slide, 3.97, 5.92, 1.32, 0.76, fill=WHITE, line=BLUE, radius=True)
    add_text(slide, "Hₜ", 4.08, 6.02, 0.35, 0.24, size=16, color=BLUE, bold=True)
    add_text(slide, "vs", 4.47, 6.02, 0.28, 0.24, size=11, color=GRAY, bold=True)
    add_text(slide, "H′ₜ", 4.78, 6.02, 0.35, 0.24, size=16, color=GREEN, bold=True)
    add_text(slide, "Same Task + Same Seed", 4.05, 6.34, 1.16, 0.20,
             size=7.7, bold=True)
    criteria = ["Task Success", "Steps / Reward", "No Regression", "Held-out Seed"]
    for i, t in enumerate(criteria):
        add_text(slide, "☑  " + t, 5.52, 5.88 + i * 0.22, 1.48, 0.20,
                 size=9, align=PP_ALIGN.LEFT, bold=True)
    add_line(slide, 7.08, 6.23, 7.42, 6.23, color=ORANGE, end=True)
    add_text(slide, "✓ ACCEPT", 7.47, 5.94, 1.52, 0.30,
             size=14, color=GREEN, bold=True)
    add_text(slide, "↶ REJECT", 7.47, 6.39, 1.52, 0.30,
             size=14, color=RED, bold=True)

    # Updated harness.
    add_box(slide, 10.83, 1.86, 2.25, 3.34, fill=LIGHT_GRAY, line=GRAY)
    add_text(slide, "Updated Gaming Harness  Hₜ₊₁", 10.93, 1.98, 2.05, 0.35,
             size=13.3, bold=True)
    add_text(slide, "validated knowledge · probes\nmemory · policies", 11.02, 2.34, 1.86, 0.44,
             size=9.3, color=GRAY)
    add_box(slide, 11.19, 2.94, 1.5, 1.15, fill=WHITE, line=BLACK, radius=True)
    add_text(slide, "✓", 11.28, 3.04, 0.38, 0.30, size=20, color=GREEN, bold=True)
    add_text(slide, "Alignment", 11.70, 3.02, 0.78, 0.21, size=8.4, bold=True)
    add_text(slide, "Recovery", 11.70, 3.31, 0.78, 0.21, size=8.4, bold=True)
    add_text(slide, "Verifier", 11.70, 3.60, 0.78, 0.21, size=8.4, bold=True)
    add_line(slide, 11.94, 4.08, 11.94, 4.34, color=BLACK, end=True)
    frozen_right = add_box(slide, 11.05, 4.35, 1.78, 0.55, fill=WHITE,
                           line=GRAY, radius=True, width=1.1)
    frozen_right.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    add_text(slide, "Frozen Gaming Model  M", 11.11, 4.44, 1.66, 0.32,
             size=11.2, bold=True)
    add_line(slide, 8.99, 6.08, 11.94, 6.08, color=GREEN, end=False)
    add_line(slide, 11.94, 6.08, 11.94, 5.20, color=GREEN, end=True)
    add_text(slide, "Accept — Update Harness", 9.17, 5.78, 1.62, 0.28,
             size=11, color=GREEN, bold=True)
    add_line(slide, 11.94, 1.86, 11.94, 0.88, color=BLACK, dash=True, end=True)
    add_text(slide, "Next Iteration", 11.12, 0.55, 1.65, 0.30,
             size=15, bold=True)

    # Reject loop back.
    add_line(slide, 7.55, 6.84, 1.31, 6.84, color=RED, end=False)
    add_line(slide, 1.31, 6.84, 1.31, 5.20, color=RED, end=True)
    add_text(slide, "Reject — No Update", 1.55, 6.88, 1.75, 0.26,
             size=11.5, color=RED, bold=True)

    prs.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
